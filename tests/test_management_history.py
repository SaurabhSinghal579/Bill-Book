import os
import tempfile
import unittest
from io import BytesIO
from urllib.parse import unquote

from openpyxl import load_workbook

import app


class ManagementHistoryTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_database = app.DATABASE
        app.DATABASE = os.path.join(self.temp_dir.name, "test.sqlite3")
        app.initialize_db()
        self.client = app.app.test_client()

    def tearDown(self):
        app.DATABASE = self.original_database
        self.temp_dir.cleanup()

    def create_bill(
        self, customer_id, item_id, quantity_id, bill_date, price, unit_id=None
    ):
        if unit_id is None:
            units = self.client.get(
                f"/api/bootstrap?date={bill_date}"
            ).json["units"]
            unit = next(
                (entry for entry in units if entry["name"] == "\u092a\u0940\u0938"),
                None,
            )
            if unit is None:
                unit = self.client.post(
                    "/api/units", json={"name": "piece"}
                ).json
            unit_id = unit["id"]
        return self.client.post(
            "/api/bills",
            json={
                "date": bill_date,
                "customer_id": customer_id,
                "items": [
                    {
                        "item_id": item_id,
                        "quantity_id": quantity_id,
                        **({"unit_id": unit_id} if unit_id is not None else {}),
                        "price": price,
                    }
                ],
            },
        )

    def test_master_crud_history_filters_and_saved_bill_snapshot(self):
        customer = self.client.post(
            "/api/customers", json={"name": "Rajendra"}
        ).json
        item = self.client.post("/api/items", json={"name": "Chawal"}).json
        quantity = self.client.post(
            "/api/quantities", json={"value": "1"}
        ).json
        bill = self.create_bill(
            customer["id"], item["id"], quantity["id"], "2026-10-11", "25"
        )
        self.assertEqual(bill.status_code, 201)

        matching = self.client.get(
            "/api/history?customer_id=1&from=2026-10-01&to=2026-10-31"
        ).json
        self.assertEqual(len(matching), 1)
        self.assertEqual(
            self.client.get("/api/history?from=2026-10-12").json, []
        )

        self.assertEqual(
            self.client.put(
                f"/api/customers/{customer['id']}", json={"name": "Ramesh"}
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.put(
                f"/api/items/{item['id']}", json={"name": "Namak"}
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.put(
                f"/api/quantities/{quantity['id']}", json={"value": "2"}
            ).status_code,
            200,
        )

        self.assertEqual(
            self.client.delete(f"/api/items/{item['id']}").status_code, 200
        )
        self.assertEqual(
            self.client.delete(f"/api/quantities/{quantity['id']}").status_code,
            200,
        )
        export = self.client.get(f"/api/bills/{bill.json['id']}/export")
        disposition = export.headers["Content-Disposition"]
        filename = unquote(disposition.split("filename*=UTF-8''", 1)[1])
        self.assertEqual(filename, "\u0930\u093e\u091c\u0947\u0902\u0926\u094d\u0930_2026-10-11.xlsx")
        sheet = load_workbook(BytesIO(export.data), data_only=True).active
        self.assertEqual(sheet["B2"].value, "\u0930\u093e\u091c\u0947\u0902\u0926\u094d\u0930")
        self.assertEqual(sheet["B5"].value, "\u091a\u093e\u0935\u0932")

        self.assertEqual(
            self.client.delete(f"/api/customers/{customer['id']}").status_code,
            200,
        )
        history = self.client.get("/api/history").json
        self.assertEqual(history[0]["customer_name"], "\u0930\u093e\u091c\u0947\u0902\u0926\u094d\u0930")
        self.assertEqual(history[0]["item_count"], 1)

    def test_bill_edit_updates_date_items_and_total_without_changing_customer(self):
        customer = self.client.post(
            "/api/customers", json={"name": "Rajendra"}
        ).json
        original_item = self.client.post(
            "/api/items", json={"name": "Chawal"}
        ).json
        new_item = self.client.post(
            "/api/items", json={"name": "Doodh"}
        ).json
        quantity = self.client.post(
            "/api/quantities", json={"value": "1"}
        ).json
        bill = self.create_bill(
            customer["id"], original_item["id"], quantity["id"], "2026-10-11", "25"
        )
        self.assertEqual(bill.status_code, 201)

        details = self.client.get(f"/api/bills/{bill.json['id']}").json
        self.assertEqual(details["items"][0]["item_name"], "\u091a\u093e\u0935\u0932")
        updated = self.client.put(
            f"/api/bills/{bill.json['id']}",
            json={
                "date": "2026-10-12",
                "items": [
                    {
                        "item_id": new_item["id"],
                        "item_name": "\u0926\u0942\u0927",
                        "quantity": "2",
                        "unit": "\u092a\u0940\u0938",
                        "price": "40",
                    }
                ],
            },
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.json["total"], 80)
        saved = self.client.get(f"/api/bills/{bill.json['id']}").json
        self.assertEqual(saved["customer_id"], customer["id"])
        self.assertEqual(saved["customer_name"], "\u0930\u093e\u091c\u0947\u0902\u0926\u094d\u0930")
        self.assertEqual(saved["date"], "2026-10-12")
        self.assertEqual(saved["items"][0]["item_name"], "\u0926\u0942\u0927")
        self.assertEqual(saved["items"][0]["quantity"], "2")
        self.assertEqual(saved["items"][0]["price"], 40)
        workbook = load_workbook(
            BytesIO(self.client.get(f"/api/bills/{bill.json['id']}/export").data),
            data_only=True,
        ).active
        self.assertEqual(workbook["B5"].value, "\u0926\u0942\u0927")
        self.assertEqual(workbook["E5"].value, 40)
        self.assertEqual(workbook["F5"].value, 80)

    def test_unit_management_and_bill_unit_snapshot(self):
        unit = self.client.post("/api/units", json={"name": "gram"})
        self.assertEqual(unit.status_code, 201)
        self.assertEqual(unit.json["name"], "\u0917\u094d\u0930\u093e\u092e")
        unit_data = unit.json

        customer = self.client.post(
            "/api/customers", json={"name": "Rajendra"}
        ).json
        item = self.client.post("/api/items", json={"name": "Chawal"}).json
        quantity = self.client.post(
            "/api/quantities", json={"value": "1"}
        ).json
        bill = self.create_bill(
            customer["id"],
            item["id"],
            quantity["id"],
            "2026-10-11",
            "25",
            unit_data["id"],
        )
        self.assertEqual(bill.status_code, 201)

        details = self.client.get(f"/api/bills/{bill.json['id']}").json
        self.assertEqual(details["items"][0]["unit"], "\u0917\u094d\u0930\u093e\u092e")
        self.assertIn(
            unit_data,
            self.client.get("/api/bootstrap?date=2026-10-11").json["units"],
        )

        updated_unit = self.client.put(
            f"/api/units/{unit_data['id']}", json={"name": "piece"}
        )
        self.assertEqual(updated_unit.status_code, 200)
        self.assertEqual(updated_unit.json["name"], "\u092a\u0940\u0938")
        updated_bill = self.client.put(
            f"/api/bills/{bill.json['id']}",
            json={
                "date": "2026-10-11",
                "items": [
                    {
                        "item_id": item["id"],
                        "quantity": "1",
                        "unit_id": unit_data["id"],
                        "price": "25",
                    }
                ],
            },
        )
        self.assertEqual(updated_bill.status_code, 200)
        self.assertEqual(
            self.client.get(f"/api/bills/{bill.json['id']}").json["items"][0]["unit"],
            "\u092a\u0940\u0938",
        )
        workbook = load_workbook(
            BytesIO(self.client.get(f"/api/bills/{bill.json['id']}/export").data),
            data_only=True,
        ).active
        self.assertEqual(workbook["D5"].value, "\u092a\u0940\u0938")

        self.assertEqual(
            self.client.delete(f"/api/units/{unit_data['id']}").status_code, 200
        )
        self.assertEqual(
            self.client.get(f"/api/bills/{bill.json['id']}").json["items"][0]["unit"],
            "\u092a\u0940\u0938",
        )

    def test_bill_uses_per_kilogram_price_for_gram_and_kilogram_units(self):
        customer = self.client.post(
            "/api/customers", json={"name": "Rajendra"}
        ).json
        item = self.client.post("/api/items", json={"name": "Chawal"}).json
        quantity = self.client.post(
            "/api/quantities", json={"value": "2"}
        ).json
        gram = self.client.post("/api/units", json={"name": "gram"}).json
        kilogram = self.client.post(
            "/api/units", json={"name": "kilogram"}
        ).json

        without_unit = self.client.post(
            "/api/bills",
            json={
                "date": "2026-10-11",
                "customer_id": customer["id"],
                "items": [
                    {
                        "item_id": item["id"],
                        "quantity_id": quantity["id"],
                        "price": "1000",
                    }
                ],
            },
        )
        self.assertEqual(without_unit.status_code, 400)

        gram_bill = self.create_bill(
            customer["id"], item["id"], quantity["id"], "2026-10-11", "1000", gram["id"]
        )
        kilogram_bill = self.create_bill(
            customer["id"], item["id"], quantity["id"], "2026-10-11", "1000", kilogram["id"]
        )
        self.assertEqual(gram_bill.status_code, 201)
        self.assertEqual(gram_bill.json["total"], 2)
        self.assertEqual(kilogram_bill.status_code, 201)
        self.assertEqual(kilogram_bill.json["total"], 2000)
        missing_unit_update = self.client.put(
            f"/api/bills/{gram_bill.json['id']}",
            json={
                "date": "2026-10-11",
                "items": [
                    {
                        "item_id": item["id"],
                        "quantity": "2",
                        "price": "1000",
                    }
                ],
            },
        )
        self.assertEqual(missing_unit_update.status_code, 400)
        gram_update = self.client.put(
            f"/api/bills/{gram_bill.json['id']}",
            json={
                "date": "2026-10-11",
                "items": [
                    {
                        "item_id": item["id"],
                        "quantity": "2",
                        "unit": "\u0917\u094d\u0930\u093e",
                        "price": "1000",
                    }
                ],
            },
        )
        self.assertEqual(gram_update.status_code, 200)
        self.assertEqual(gram_update.json["total"], 2)

    def test_delete_bill_removes_bill_and_its_items(self):
        customer = self.client.post(
            "/api/customers", json={"name": "Rajendra"}
        ).json
        item = self.client.post("/api/items", json={"name": "Chawal"}).json
        quantity = self.client.post(
            "/api/quantities", json={"value": "1"}
        ).json
        bill = self.create_bill(
            customer["id"], item["id"], quantity["id"], "2026-10-11", "25"
        )
        bill_id = bill.json["id"]

        deleted = self.client.delete(f"/api/bills/{bill_id}")
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(deleted.json, {"deleted": bill_id})
        self.assertEqual(self.client.get("/api/history").json, [])
        self.assertEqual(self.client.get(f"/api/bills/{bill_id}").status_code, 404)
        self.assertEqual(
            self.client.get(f"/api/bills/{bill_id}/export").status_code, 404
        )
        self.assertEqual(self.client.delete(f"/api/bills/{bill_id}").status_code, 404)

    def test_master_names_keep_english_and_hindi_for_inline_edit(self):
        customer = self.client.post(
            "/api/customers",
            json={"english_name": "Rajendra", "name": "\u0930\u093e\u091c\u0947\u0902\u0926\u094d\u0930"},
        ).json
        item = self.client.post(
            "/api/items",
            json={"english_name": "Chawal", "name": "\u091a\u093e\u0935\u0932"},
        ).json
        unit = self.client.post(
            "/api/units",
            json={"english_name": "gram", "name": "\u0917\u094d\u0930\u093e\u092e"},
        ).json

        data = self.client.get("/api/bootstrap?date=2026-10-11").json
        self.assertEqual(
            next(entry for entry in data["customers"] if entry["id"] == customer["id"])[
                "english_name"
            ],
            "Rajendra",
        )
        self.assertEqual(
            next(entry for entry in data["items"] if entry["id"] == item["id"])[
                "english_name"
            ],
            "Chawal",
        )
        self.assertEqual(
            next(entry for entry in data["units"] if entry["id"] == unit["id"])[
                "english_name"
            ],
            "gram",
        )

        updates = (
            ("/api/customers", customer["id"], "Ramesh", "\u0930\u092e\u0947\u0936"),
            ("/api/items", item["id"], "Rice", "\u091a\u093e\u0935\u0932"),
            ("/api/units", unit["id"], "kilogram", "\u0915\u093f\u0932\u094b\u0917\u094d\u0930\u093e\u092e"),
        )
        for endpoint, record_id, english_name, hindi in updates:
            with self.subTest(endpoint=endpoint):
                response = self.client.put(
                    f"{endpoint}/{record_id}",
                    json={"english_name": english_name, "name": hindi},
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json["english_name"], english_name)
                self.assertEqual(response.json["name"], hindi)

    def test_legacy_master_names_get_romanized_for_editing(self):
        with app.connect_db() as connection:
            connection.execute(
                "INSERT INTO customers(name, english_name) VALUES (?, '')",
                ("\u0930\u093e\u091c\u0947\u0902\u0926\u094d\u0930",),
            )
            connection.execute(
                "INSERT INTO items(name, unit, english_name) VALUES (?, '', '')",
                ("\u091a\u093e\u0935\u0932",),
            )
            connection.execute(
                "INSERT INTO units(name, english_name) VALUES (?, '')",
                ("\u0917\u094d\u0930\u093e\u092e",),
            )

        app.initialize_db()
        data = self.client.get("/api/bootstrap?date=2026-10-11").json
        self.assertTrue(data["customers"][0]["english_name"])
        self.assertTrue(data["items"][0]["english_name"])
        self.assertTrue(data["units"][0]["english_name"])
        self.assertTrue(data["customers"][0]["english_name"].isascii())


if __name__ == "__main__":
    unittest.main()
