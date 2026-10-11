from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import os
import re
import sqlite3

from flask import Flask, jsonify, render_template, request, send_file
from indic_transliteration import sanscript
from indic_transliteration.sanscript import transliterate
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


app = Flask(__name__)
DATABASE = os.path.join(app.instance_path, "bills.sqlite3")
os.makedirs(app.instance_path, exist_ok=True)


class ClosingConnection(sqlite3.Connection):
    def __exit__(self, exc_type, exc_value, traceback):
        try:
            return super().__exit__(exc_type, exc_value, traceback)
        finally:
            self.close()


def connect_db():
    connection = sqlite3.connect(DATABASE, factory=ClosingConnection)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def romanized_name(text):
    value = str(text).strip()
    if not value or re.search(r"[A-Za-z]", value):
        return value
    return transliterate(value, sanscript.DEVANAGARI, sanscript.ITRANS).lower()


def initialize_db():
    with connect_db() as connection:
        connection.execute("PRAGMA journal_mode = WAL")
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS customers (
                id INTEGER PRIMARY KEY,
                name TEXT NOT NULL UNIQUE,
                english_name TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS items (
                id INTEGER PRIMARY KEY,
                name TEXT NOT NULL UNIQUE,
                unit TEXT NOT NULL DEFAULT '',
                english_name TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS item_prices (
                id INTEGER PRIMARY KEY,
                item_id INTEGER NOT NULL REFERENCES items(id) ON DELETE CASCADE,
                price_paise INTEGER NOT NULL CHECK (price_paise >= 0),
                UNIQUE(item_id, price_paise)
            );
            CREATE TABLE IF NOT EXISTS quantities (
                id INTEGER PRIMARY KEY,
                value TEXT NOT NULL UNIQUE
            );
            CREATE TABLE IF NOT EXISTS units (
                id INTEGER PRIMARY KEY,
                name TEXT NOT NULL UNIQUE,
                english_name TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS bills (
                id INTEGER PRIMARY KEY,
                bill_date TEXT NOT NULL,
                customer_id INTEGER REFERENCES customers(id),
                customer_name TEXT NOT NULL,
                total_paise INTEGER NOT NULL,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS bill_items (
                id INTEGER PRIMARY KEY,
                bill_id INTEGER NOT NULL REFERENCES bills(id) ON DELETE CASCADE,
                item_id INTEGER REFERENCES items(id),
                item_name TEXT NOT NULL,
                quantity TEXT NOT NULL,
                unit TEXT NOT NULL,
                unit_price_paise INTEGER NOT NULL,
                line_total_paise INTEGER NOT NULL
            );
            CREATE INDEX IF NOT EXISTS bills_by_date ON bills(bill_date, id DESC);
            """
        )
        for table in ("customers", "items", "units"):
            columns = {
                row["name"]
                for row in connection.execute(f"PRAGMA table_info({table})")
            }
            if "english_name" not in columns:
                connection.execute(
                    f"ALTER TABLE {table} ADD COLUMN english_name TEXT NOT NULL DEFAULT ''"
                )
            records = connection.execute(
                f"SELECT id, name FROM {table} WHERE english_name = ''"
            ).fetchall()
            for record in records:
                connection.execute(
                    f"UPDATE {table} SET english_name = ? WHERE id = ?",
                    (romanized_name(record["name"]), record["id"]),
                )


def json_error(message, status=400):
    return jsonify({"error": message}), status


def money_to_paise(value):
    amount = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return int(amount * 100)


def bill_line_total_paise(quantity, price_paise, unit):
    unit_name = re.sub(r"[\s._-]+", "", str(unit).strip().casefold())
    gram_units = {
        "g", "gm", "gms", "gram", "grams",
        "ग्राम", "ग्राम्स", "ग्रा", "ग्र",
    }
    unit_factor = Decimal("0.001") if unit_name in gram_units else Decimal("1")
    return int(
        (quantity * price_paise * unit_factor).quantize(
            Decimal("1"), rounding=ROUND_HALF_UP
        )
    )


def hindi_name(text):
    value = str(text).strip()
    if not value or any("\u0900" <= character <= "\u097f" for character in value):
        return value
    common_spellings = {
        "rajendra": "\u0930\u093e\u091c\u0947\u0902\u0926\u094d\u0930",
        "chawal": "\u091a\u093e\u0935\u0932",
        "doodh": "\u0926\u0942\u0927",
        "namak": "\u0928\u092e\u0915",
        "g": "\u0917\u094d\u0930\u093e\u092e",
        "gm": "\u0917\u094d\u0930\u093e\u092e",
        "gms": "\u0917\u094d\u0930\u093e\u092e",
        "gram": "\u0917\u094d\u0930\u093e\u092e",
        "grams": "\u0917\u094d\u0930\u093e\u092e",
        "kg": "\u0915\u093f\u0917\u094d\u0930\u093e",
        "kgs": "\u0915\u093f\u0917\u094d\u0930\u093e",
        "kilogram": "\u0915\u093f\u0932\u094b\u0917\u094d\u0930\u093e\u092e",
        "kilograms": "\u0915\u093f\u0932\u094b\u0917\u094d\u0930\u093e\u092e",
        "piece": "\u092a\u0940\u0938",
        "pieces": "\u092a\u0940\u0938",
        "pcs": "\u092a\u0940\u0938",
    }

    def convert_word(match):
        word = match.group(0)
        normalized = word.lower()
        if normalized in common_spellings:
            return common_spellings[normalized]
        normalized = re.sub(r"chh", "Ch", normalized)
        normalized = re.sub(r"aa", "A", normalized)
        normalized = re.sub(r"ee|ii", "I", normalized)
        normalized = re.sub(r"oo|uu", "U", normalized)
        normalized = re.sub(r"n(?=[kgcjtdpb])", "M", normalized)
        normalized = normalized.replace("w", "v")
        converted = transliterate(
            normalized, sanscript.ITRANS, sanscript.DEVANAGARI
        )
        return converted.removesuffix("्")

    return re.sub(r"[A-Za-z]+", convert_word, value)


def bill_rows_for_date(connection, bill_date):
    return connection.execute(
        """
        SELECT id, bill_date, customer_name, total_paise, created_at
        FROM bills WHERE bill_date = ? ORDER BY id DESC
        """,
        (bill_date,),
    ).fetchall()


def bill_summary(rows):
    return [
        {
            "id": row["id"],
            "date": row["bill_date"],
            "customer_name": row["customer_name"],
            "total": row["total_paise"] / 100,
            "created_at": row["created_at"],
        }
        for row in rows
    ]


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/bootstrap")
def bootstrap():
    selected_date = request.args.get("date", date.today().isoformat())
    try:
        date.fromisoformat(selected_date)
    except ValueError:
        return json_error("Date format sahi nahi hai.")

    with connect_db() as connection:
        customers = connection.execute(
            "SELECT id, name, english_name FROM customers ORDER BY name COLLATE NOCASE"
        ).fetchall()
        items = connection.execute(
            "SELECT id, name, english_name, unit FROM items ORDER BY name COLLATE NOCASE"
        ).fetchall()
        prices = connection.execute(
            "SELECT id, item_id, price_paise FROM item_prices ORDER BY price_paise"
        ).fetchall()
        quantities = connection.execute(
            "SELECT id, value FROM quantities ORDER BY CAST(value AS REAL)"
        ).fetchall()
        units = connection.execute(
            "SELECT id, name, english_name FROM units ORDER BY name COLLATE NOCASE"
        ).fetchall()
        bills = bill_rows_for_date(connection, selected_date)

    price_map = {}
    for price in prices:
        price_map.setdefault(price["item_id"], []).append(
            {"id": price["id"], "amount": price["price_paise"] / 100}
        )

    return jsonify(
        {
            "customers": [dict(row) for row in customers],
            "items": [
                {
                    **dict(row),
                    "prices": price_map.get(row["id"], []),
                }
                for row in items
            ],
            "quantities": [dict(row) for row in quantities],
            "units": [dict(row) for row in units],
            "bills": bill_summary(bills),
            "date": selected_date,
        }
    )


@app.get("/api/history")
def bill_history():
    conditions = []
    parameters = []
    customer_id = request.args.get("customer_id", "").strip()
    if customer_id:
        try:
            parameters.append(int(customer_id))
        except ValueError:
            return json_error("Customer filter sahi nahi hai.")
        conditions.append("b.customer_id = ?")

    for parameter, operator in (("from", ">="), ("to", "<=")):
        selected_date = request.args.get(parameter, "").strip()
        if selected_date:
            try:
                date.fromisoformat(selected_date)
            except ValueError:
                return json_error("Date format sahi nahi hai.")
            conditions.append(f"b.bill_date {operator} ?")
            parameters.append(selected_date)

    where_clause = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    with connect_db() as connection:
        rows = connection.execute(
            f"""SELECT b.id, b.bill_date, b.customer_id, b.customer_name,
                b.total_paise, b.created_at, COUNT(bi.id) AS item_count
            FROM bills AS b LEFT JOIN bill_items AS bi ON bi.bill_id = b.id
            {where_clause}
            GROUP BY b.id ORDER BY b.bill_date DESC, b.id DESC""",
            parameters,
        ).fetchall()
    return jsonify(
        [
            {
                "id": row["id"],
                "date": row["bill_date"],
                "customer_id": row["customer_id"],
                "customer_name": row["customer_name"],
                "total": row["total_paise"] / 100,
                "item_count": row["item_count"],
                "created_at": row["created_at"],
            }
            for row in rows
        ]
    )


@app.post("/api/customers")
def create_customer():
    data = request.get_json(silent=True) or {}
    name = hindi_name(data.get("name", ""))
    english = str(data.get("english_name", "")).strip() or romanized_name(
        data.get("name", "")
    )
    if not name or not english or len(english) > 120:
        return json_error("Customer ka naam likhiye.")
    try:
        with connect_db() as connection:
            cursor = connection.execute(
                "INSERT INTO customers(name, english_name) VALUES (?, ?)",
                (name, english),
            )
        return jsonify(
            {"id": cursor.lastrowid, "name": name, "english_name": english}
        ), 201
    except sqlite3.IntegrityError:
        return json_error("Yeh customer pehle se maujood hai.", 409)


@app.put("/api/customers/<int:customer_id>")
def update_customer(customer_id):
    data = request.get_json(silent=True) or {}
    name = hindi_name(data.get("name", ""))
    english = str(data.get("english_name", "")).strip() or romanized_name(
        data.get("name", "")
    )
    if not name or not english or len(english) > 120:
        return json_error("Customer ka naam likhiye.")
    try:
        with connect_db() as connection:
            cursor = connection.execute(
                "UPDATE customers SET name = ?, english_name = ? WHERE id = ?",
                (name, english, customer_id),
            )
        if not cursor.rowcount:
            return json_error("Customer nahi mila.", 404)
        return jsonify(
            {"id": customer_id, "name": name, "english_name": english}
        )
    except sqlite3.IntegrityError:
        return json_error("Yeh customer pehle se maujood hai.", 409)


@app.delete("/api/customers/<int:customer_id>")
def delete_customer(customer_id):
    with connect_db() as connection:
        cursor = connection.execute(
            "UPDATE bills SET customer_id = NULL WHERE customer_id = ?",
            (customer_id,),
        )
        deleted = connection.execute(
            "DELETE FROM customers WHERE id = ?", (customer_id,)
        )
    if not deleted.rowcount:
        return json_error("Customer nahi mila.", 404)
    return jsonify({"deleted": customer_id})


@app.post("/api/items")
def create_item():
    data = request.get_json(silent=True) or {}
    name = hindi_name(data.get("name", ""))
    english = str(data.get("english_name", "")).strip() or romanized_name(
        data.get("name", "")
    )
    if not name or not english or len(english) > 120:
        return json_error("Item ka naam likhiye.")

    try:
        with connect_db() as connection:
            cursor = connection.execute(
                "INSERT INTO items(name, english_name) VALUES (?, ?)",
                (name, english),
            )
            item_id = cursor.lastrowid
        return jsonify(
            {"id": item_id, "name": name, "english_name": english}
        ), 201
    except sqlite3.IntegrityError:
        return json_error("Yeh item pehle se maujood hai.", 409)


@app.put("/api/items/<int:item_id>")
def update_item(item_id):
    data = request.get_json(silent=True) or {}
    name = hindi_name(data.get("name", ""))
    english = str(data.get("english_name", "")).strip() or romanized_name(
        data.get("name", "")
    )
    if not name or not english or len(english) > 120:
        return json_error("Item ka naam likhiye.")
    try:
        with connect_db() as connection:
            cursor = connection.execute(
                "UPDATE items SET name = ?, english_name = ? WHERE id = ?",
                (name, english, item_id),
            )
        if not cursor.rowcount:
            return json_error("Item nahi mila.", 404)
        return jsonify(
            {"id": item_id, "name": name, "english_name": english}
        )
    except sqlite3.IntegrityError:
        return json_error("Yeh item pehle se maujood hai.", 409)


@app.delete("/api/items/<int:item_id>")
def delete_item(item_id):
    with connect_db() as connection:
        connection.execute(
            "UPDATE bill_items SET item_id = NULL WHERE item_id = ?", (item_id,)
        )
        deleted = connection.execute("DELETE FROM items WHERE id = ?", (item_id,))
    if not deleted.rowcount:
        return json_error("Item nahi mila.", 404)
    return jsonify({"deleted": item_id})


@app.post("/api/transliterate")
def transliterate_name():
    data = request.get_json(silent=True) or {}
    return jsonify({"hindi": hindi_name(data.get("text", ""))})


@app.post("/api/quantities")
def create_quantity():
    data = request.get_json(silent=True) or {}
    try:
        value = Decimal(str(data.get("value", "")))
        if not value.is_finite() or value <= 0:
            raise InvalidOperation
        normalized = format(value.normalize(), "f")
    except (InvalidOperation, ValueError):
        return json_error("Quantity zero se bada number hona chahiye.")
    try:
        with connect_db() as connection:
            cursor = connection.execute(
                "INSERT INTO quantities(value) VALUES (?)", (normalized,)
            )
        return jsonify({"id": cursor.lastrowid, "value": normalized}), 201
    except sqlite3.IntegrityError:
        return json_error("Yeh quantity pehle se maujood hai.", 409)


@app.put("/api/quantities/<int:quantity_id>")
def update_quantity(quantity_id):
    data = request.get_json(silent=True) or {}
    try:
        value = Decimal(str(data.get("value", "")))
        if not value.is_finite() or value <= 0:
            raise InvalidOperation
        normalized = format(value.normalize(), "f")
    except (InvalidOperation, ValueError):
        return json_error("Quantity zero se bada number hona chahiye.")
    try:
        with connect_db() as connection:
            cursor = connection.execute(
                "UPDATE quantities SET value = ? WHERE id = ?",
                (normalized, quantity_id),
            )
        if not cursor.rowcount:
            return json_error("Quantity nahi mili.", 404)
        return jsonify({"id": quantity_id, "value": normalized})
    except sqlite3.IntegrityError:
        return json_error("Yeh quantity pehle se maujood hai.", 409)


@app.delete("/api/quantities/<int:quantity_id>")
def delete_quantity(quantity_id):
    with connect_db() as connection:
        deleted = connection.execute(
            "DELETE FROM quantities WHERE id = ?", (quantity_id,)
        )
    if not deleted.rowcount:
        return json_error("Quantity nahi mili.", 404)
    return jsonify({"deleted": quantity_id})


@app.post("/api/units")
def create_unit():
    data = request.get_json(silent=True) or {}
    name = hindi_name(data.get("name", ""))
    english = str(data.get("english_name", "")).strip() or romanized_name(
        data.get("name", "")
    )
    if not name or not english or len(name) > 30 or len(english) > 30:
        return json_error("Unit ka naam likhiye.")
    try:
        with connect_db() as connection:
            cursor = connection.execute(
                "INSERT INTO units(name, english_name) VALUES (?, ?)",
                (name, english),
            )
        return jsonify(
            {"id": cursor.lastrowid, "name": name, "english_name": english}
        ), 201
    except sqlite3.IntegrityError:
        return json_error("Yeh unit pehle se maujood hai.", 409)


@app.put("/api/units/<int:unit_id>")
def update_unit(unit_id):
    data = request.get_json(silent=True) or {}
    name = hindi_name(data.get("name", ""))
    english = str(data.get("english_name", "")).strip() or romanized_name(
        data.get("name", "")
    )
    if not name or not english or len(name) > 30 or len(english) > 30:
        return json_error("Unit ka naam likhiye.")
    try:
        with connect_db() as connection:
            cursor = connection.execute(
                "UPDATE units SET name = ?, english_name = ? WHERE id = ?",
                (name, english, unit_id),
            )
        if not cursor.rowcount:
            return json_error("Unit nahi mila.", 404)
        return jsonify(
            {"id": unit_id, "name": name, "english_name": english}
        )
    except sqlite3.IntegrityError:
        return json_error("Yeh unit pehle se maujood hai.", 409)


@app.delete("/api/units/<int:unit_id>")
def delete_unit(unit_id):
    with connect_db() as connection:
        deleted = connection.execute("DELETE FROM units WHERE id = ?", (unit_id,))
    if not deleted.rowcount:
        return json_error("Unit nahi mila.", 404)
    return jsonify({"deleted": unit_id})


@app.get("/api/bills/<int:bill_id>")
def get_bill(bill_id):
    with connect_db() as connection:
        bill = connection.execute(
            "SELECT id, bill_date, customer_id, customer_name FROM bills WHERE id = ?",
            (bill_id,),
        ).fetchone()
        if bill is None:
            return json_error("Bill nahi mila.", 404)
        rows = connection.execute(
            """SELECT item_id, item_name, quantity, unit, unit_price_paise
            FROM bill_items WHERE bill_id = ? ORDER BY id""",
            (bill_id,),
        ).fetchall()
    return jsonify(
        {
            "id": bill["id"],
            "date": bill["bill_date"],
            "customer_id": bill["customer_id"],
            "customer_name": bill["customer_name"],
            "items": [
                {
                    "item_id": row["item_id"],
                    "item_name": row["item_name"],
                    "quantity": row["quantity"],
                    "unit": row["unit"],
                    "price": row["unit_price_paise"] / 100,
                }
                for row in rows
            ],
        }
    )


@app.delete("/api/bills/<int:bill_id>")
def delete_bill(bill_id):
    with connect_db() as connection:
        deleted = connection.execute("DELETE FROM bills WHERE id = ?", (bill_id,))
    if not deleted.rowcount:
        return json_error("Bill nahi mila.", 404)
    return jsonify({"deleted": bill_id})


@app.put("/api/bills/<int:bill_id>")
def update_bill(bill_id):
    data = request.get_json(silent=True) or {}
    bill_date = str(data.get("date", ""))
    requested_rows = data.get("items")
    try:
        date.fromisoformat(bill_date)
        if not isinstance(requested_rows, list) or not requested_rows:
            raise ValueError
    except (TypeError, ValueError):
        return json_error("Date aur kam se kam ek item zaroori hai.")

    try:
        with connect_db() as connection:
            bill = connection.execute(
                "SELECT id FROM bills WHERE id = ?", (bill_id,)
            ).fetchone()
            if bill is None:
                return json_error("Bill nahi mila.", 404)

            calculated_rows = []
            total_paise = 0
            for requested in requested_rows:
                raw_item_id = requested.get("item_id")
                item_id = int(raw_item_id) if raw_item_id not in (None, "") else None
                item = None
                if item_id is not None:
                    item = connection.execute(
                        "SELECT id, name FROM items WHERE id = ?", (item_id,)
                    ).fetchone()
                    if item is None:
                        return json_error("Bill mein item valid nahi hai.")
                item_name = item["name"] if item else hindi_name(requested.get("item_name", ""))
                if not item_name:
                    return json_error("Har row mein item chunein.")

                unit = str(requested.get("unit", "")).strip()
                raw_unit_id = requested.get("unit_id")
                if raw_unit_id not in (None, ""):
                    selected_unit = connection.execute(
                        "SELECT name FROM units WHERE id = ?", (int(raw_unit_id),)
                    ).fetchone()
                    if selected_unit is None:
                        return json_error("Bill mein unit valid nahi hai.")
                    unit = selected_unit["name"]
                if len(unit) > 30:
                    return json_error("Har item ke liye valid unit chunein.")
                if not unit:
                    return json_error("Har bill row mein unit chunein.")

                quantity_value = Decimal(str(requested["quantity"]))
                price_paise = money_to_paise(requested["price"])
                if not quantity_value.is_finite() or quantity_value <= 0 or price_paise < 0:
                    return json_error("Har row mein valid quantity aur price dijiye.")
                line_total = bill_line_total_paise(quantity_value, price_paise, unit)
                calculated_rows.append(
                    (
                        item_id,
                        item_name,
                        format(quantity_value.normalize(), "f"),
                        unit,
                        price_paise,
                        line_total,
                    )
                )
                total_paise += line_total

            connection.execute(
                "UPDATE bills SET bill_date = ?, total_paise = ? WHERE id = ?",
                (bill_date, total_paise, bill_id),
            )
            connection.execute("DELETE FROM bill_items WHERE bill_id = ?", (bill_id,))
            connection.executemany(
                """INSERT INTO bill_items(
                    bill_id, item_id, item_name, quantity, unit,
                    unit_price_paise, line_total_paise
                ) VALUES (?, ?, ?, ?, ?, ?, ?)""",
                [(bill_id, *row) for row in calculated_rows],
            )
        return jsonify({"id": bill_id, "total": total_paise / 100}), 200
    except (KeyError, TypeError, ValueError, InvalidOperation):
        return json_error("Bill details sahi format mein nahi hain.")


@app.post("/api/bills")
def create_bill():
    data = request.get_json(silent=True) or {}
    bill_date = str(data.get("date", ""))
    try:
        date.fromisoformat(bill_date)
        customer_id = int(data.get("customer_id"))
        requested_rows = data.get("items")
        if not isinstance(requested_rows, list) or not requested_rows:
            raise ValueError
    except (TypeError, ValueError):
        return json_error("Date, customer aur kam se kam ek item zaroori hai.")

    try:
        with connect_db() as connection:
            customer = connection.execute(
                "SELECT id, name FROM customers WHERE id = ?", (customer_id,)
            ).fetchone()
            if customer is None:
                return json_error("Customer nahi mila. Dropdown refresh karke dobara try karein.", 404)

            calculated_rows = []
            total_paise = 0
            for requested in requested_rows:
                item_id = int(requested["item_id"])
                quantity_id = int(requested["quantity_id"])
                raw_unit_id = requested.get("unit_id")
                unit = ""
                if raw_unit_id not in (None, ""):
                    selected_unit = connection.execute(
                        "SELECT name FROM units WHERE id = ?", (int(raw_unit_id),)
                    ).fetchone()
                    if selected_unit is None:
                        return json_error("Bill mein unit valid nahi hai.")
                    unit = selected_unit["name"]
                if not unit:
                    return json_error("Har bill row mein unit chunein.")
                price_paise = money_to_paise(requested["price"])
                if len(unit) > 30 or price_paise < 0:
                    return json_error("Har item ke liye valid price dijiye.")
                item = connection.execute(
                    "SELECT id, name FROM items WHERE id = ?", (item_id,)
                ).fetchone()
                quantity = connection.execute(
                    "SELECT value FROM quantities WHERE id = ?", (quantity_id,)
                ).fetchone()
                if item is None or quantity is None:
                    return json_error("Bill mein koi item, price ya quantity valid nahi hai.", 400)
                line_total = bill_line_total_paise(
                    Decimal(quantity["value"]), price_paise, unit
                )
                calculated_rows.append(
                    (item_id, item["name"], quantity["value"], unit, price_paise, line_total)
                )
                total_paise += line_total

            cursor = connection.execute(
                "INSERT INTO bills(bill_date, customer_id, customer_name, total_paise) VALUES (?, ?, ?, ?)",
                (bill_date, customer_id, customer["name"], total_paise),
            )
            bill_id = cursor.lastrowid
            connection.executemany(
                """
                INSERT INTO bill_items(
                    bill_id, item_id, item_name, quantity, unit,
                    unit_price_paise, line_total_paise
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                [(bill_id, *row) for row in calculated_rows],
            )
        return jsonify({"id": bill_id, "total": total_paise / 100}), 201
    except (KeyError, TypeError, ValueError, InvalidOperation):
        return json_error("Bill details sahi format mein nahi hain.")


@app.get("/api/bills/<int:bill_id>/export")
def export_bill(bill_id):
    with connect_db() as connection:
        bill = connection.execute(
            "SELECT * FROM bills WHERE id = ?", (bill_id,)
        ).fetchone()
        rows = connection.execute(
            """SELECT item_name, quantity, unit, unit_price_paise, line_total_paise
            FROM bill_items WHERE bill_id = ? ORDER BY id""",
            (bill_id,),
        ).fetchall()
    if bill is None:
        return json_error("Bill nahi mila.", 404)
    customer_filename = re.sub(
        r'[<>:"/\\|?*\x00-\x1f]', "_", bill["customer_name"]
    ).strip(" .") or "customer"
    return make_workbook(
        bill=bill,
        rows=rows,
        filename=f"{customer_filename}_{bill['bill_date']}.xlsx",
    )


def make_workbook(bill, rows, filename):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bill"
    sheet.append(["Bill No.", bill["id"], "Date", bill["bill_date"]])
    sheet.append(["Customer", bill["customer_name"]])
    sheet.append([])
    sheet.append(["S.No.", "Item", "Qty", "Unit", "Price", "Amount"])
    for header in sheet[4]:
        header.font = Font(bold=True, color="FFFFFF")
        header.fill = PatternFill("solid", fgColor="23594F")
    for serial_number, row in enumerate(rows, start=1):
        sheet.append(
            [
                serial_number,
                row["item_name"],
                float(row["quantity"]),
                row["unit"],
                row["unit_price_paise"] / 100,
                row["line_total_paise"] / 100,
            ]
        )
    total_row = sheet.max_row + 1
    sheet.cell(total_row, 5, "Total")
    sheet.cell(total_row, 6, bill["total_paise"] / 100)
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=5, max_row=total_row, min_col=5, max_col=6):
        for cell in row:
            if cell.column == 6 or (cell.column == 5 and cell.row < total_row):
                cell.number_format = '"₹"#,##0.00'
    widths = [12, 28, 12, 18, 16, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A5"
    if rows:
        sheet.auto_filter.ref = f"A4:F{total_row - 1}"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


initialize_db()

if __name__ == "__main__":
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1", port=5000)
